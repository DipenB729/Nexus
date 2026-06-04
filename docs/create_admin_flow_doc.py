from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("Admin-Flow-CRUD-Documentation.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(widths[idx] * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
        set_cell_shading(hdr[i], "E8EEF5")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    set_table_geometry(table, widths)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def status_rows():
    return [
        ("Authentication", "Login as admin@nexus.local", "PASS", "Admin session opened /admin/dashboard after backend restart."),
        ("Page routing", "40 admin routes", "PASS", "All tested routes rendered authenticated admin content."),
        ("Dashboard", "Read operational summaries", "PASS", "Dashboard and inventory summaries returned data."),
        ("Patients", "Create patient", "FAIL", "POST /api/admin/patients returns 500. Backend log: MongoDB EF provider does not support transactions in AdminPatientsController.Create."),
        ("Patients", "Read registry", "PASS", "GET /api/admin/patients returned seeded patient records."),
        ("Patients", "Update/deactivate", "BLOCKED", "Blocked because patient creation failed; existing patient mutation was avoided."),
        ("Services", "Create, update, delete", "PASS", "Hospital service supports true DELETE and was cleaned up."),
        ("Appointments", "Read queue and update token settings", "PASS", "Queue and token settings endpoints returned successfully."),
        ("Admissions", "Create/discharge lifecycle", "BLOCKED", "Blocked by patient create failure; admissions list read passed."),
        ("Master Setup - Departments", "Create, update, deactivate", "PASS", "Status endpoint used for delete/deactivate behavior."),
        ("Master Setup - Categories", "Create, update, deactivate", "PASS", "Status endpoint used for delete/deactivate behavior."),
        ("Master Setup - Wards", "Create, update", "PASS", "Created against an existing branch/department."),
        ("Master Setup - Beds", "Create, update", "PASS", "Created against the test ward."),
        ("Master Setup - Doctors", "Create, update/deactivate", "PASS", "Doctor master workflow passed."),
        ("Master Setup - Staff", "Create, update/deactivate", "PASS", "Staff master workflow passed."),
        ("Billing - Charges", "Create, update/deactivate", "PASS", "Write path passed; page/list path can still hit Mongo projection errors."),
        ("Billing - Payment Methods", "Create, update/deactivate", "PASS", "Write path passed; page/list path can still hit Mongo projection errors."),
        ("Billing - Partners", "Create, update", "PASS", "Partner workflow passed."),
        ("Billing - Rules", "Create, update/deactivate", "PASS", "Rule workflow passed."),
        ("Billing - Bills", "Create, update, payment", "BLOCKED", "Blocked by patient create failure."),
        ("Inventory - Units", "Create, update/deactivate", "PASS", "Inventory unit workflow passed."),
        ("Inventory - Categories", "Create, update/deactivate", "PASS", "Inventory category workflow passed."),
        ("Inventory - Suppliers", "Create, update/deactivate", "PASS", "Supplier workflow passed."),
        ("Inventory - Medicines", "Create/update", "FAIL", "Mongo LINQ translation failure in inventory service after save/read mapping."),
        ("Inventory - Items", "Create/update", "FAIL", "Mongo LINQ translation failure in inventory service after save/read mapping."),
        ("Inventory - Locations", "Create/update", "FAIL", "Mongo LINQ translation failure in inventory service after save/read mapping."),
        ("Inventory - Batches", "Read", "PASS", "Batch listing returned data."),
        ("Inventory - Purchase Orders", "Create/update", "BLOCKED", "Blocked by stock location creation failure."),
        ("Inventory - Purchase Invoices", "Read", "FAIL", "Mongo cannot translate Include/OrderBy chain in EfInventoryAdminService.GetPurchaseInvoicesAsync."),
        ("Inventory - Purchase Returns", "Read", "FAIL", "Mongo cannot translate Include/OrderBy chain in EfInventoryAdminService.GetPurchaseReturnsAsync."),
        ("Inventory - Supplier Dues", "Read", "PASS", "Supplier dues summary returned data."),
        ("Inventory - Transfers", "Read", "FAIL", "Mongo cannot translate Include/OrderBy chain in EfInventoryAdminService.GetTransfersAsync."),
        ("Inventory - Adjustments", "Create/update", "BLOCKED", "Blocked by stock location creation failure."),
        ("Laboratory - Lab Tests", "Create, update/deactivate", "PASS", "Lab test workflow passed."),
        ("Laboratory - Packages", "Create, update/deactivate", "PASS", "Package workflow passed."),
        ("Monitoring", "Read reports and audit", "PASS", "Reports dashboard and audit log returned data."),
        ("Roles & Access", "Create role, update/deactivate, read profile", "PASS", "Role access workflow passed."),
        ("Settings - Organization", "Read/update", "PASS", "Organization GET/PUT passed."),
        ("Settings - Notifications/System/Security", "Read/update", "PARTIAL", "UI pages load via control center; standalone GET endpoints are not exposed, only PUT endpoints."),
        ("Notifications", "Read notification center", "PASS", "Notification list returned data."),
    ]


def main():
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Nexus Admin Flow and CRUD Test Documentation")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("0B2545")

    doc.add_paragraph(f"Tested on {datetime.now().strftime('%Y-%m-%d %H:%M')} local time. Environment: Angular admin client on 127.0.0.1:4201, .NET API on localhost:5180, MongoDB-backed EF provider.")
    doc.add_paragraph("Scope: admin workspace routing, page rendering, and CRUD/state-change checks across dashboard, patients, services, appointments, admissions, master setup, billing, inventory, laboratory, monitoring, roles/access, settings, and notifications.")

    doc.add_heading("Executive Summary", level=1)
    add_bullets(doc, [
        "Admin login works after restarting the stale backend process.",
        "All tested admin routes rendered authenticated admin content.",
        "50 of 73 CRUD/state-change API checks passed in the main pass.",
        "Most admin modules do not expose hard DELETE. The tested delete equivalent is deactivate/status change where available.",
        "Main blockers are backend/provider compatibility issues: unsupported Mongo EF transactions and untranslatable LINQ projections/includes.",
    ])

    doc.add_heading("Admin Flow Structure", level=1)
    add_table(
        doc,
        ["Flow Area", "Routes / Sections", "Observed Structure"],
        [
            ("Entry", "/auth/login -> /admin/dashboard", "Role-based login stores JWT session and opens the dashboard shell."),
            ("Shell", "app-navbar + top navbar + router outlet", "Sidebar controls admin navigation groups; header shows workspace, notifications, and profile menu."),
            ("Overview", "/admin/dashboard", "Hospital command dashboard with beds, patients, appointments, sales, stock, billing, and occupancy summaries."),
            ("Core Ops", "/admin/patients, /admin/bookings, /admin/admissions, /admin/notifications", "Patient registry, appointment control, admission control, and notification center."),
            ("Master Setup", "/admin/masters/:section", "Departments, doctors, staff, patient categories, wards, and beds."),
            ("Billing", "/admin/billing/:section", "Bills, charges, payment methods, insurance/corporate partners, and billing rules."),
            ("Supply Chain", "/admin/inventory/:section", "Dashboard, units, categories, medicines, items, suppliers, locations, batches, purchase flow, supplier dues, transfers, and adjustments."),
            ("Lab & Packages", "/admin/laboratory/:section", "Lab tests and package catalogs."),
            ("Monitoring", "/admin/monitoring/reports, /admin/monitoring/audit", "Reports dashboard and audit trail."),
            ("Administration", "/admin/roles, /admin/settings/:section", "Role permissions, admin accounts, organization, notifications, system, backups, and security settings."),
        ],
        [1.3, 2.1, 3.1],
    )

    doc.add_heading("Page Rendering Coverage", level=1)
    add_table(
        doc,
        ["Group", "Routes Tested", "Result"],
        [
            ("Overview", "/admin/dashboard", "PASS"),
            ("Core Ops", "/admin/patients, /admin/services, /admin/bookings, /admin/admissions, /admin/notifications", "PASS"),
            ("Master Setup", "/admin/masters/departments, patient-categories, wards, beds, doctors, staff", "PASS"),
            ("Billing", "/admin/billing/bills, charges, payment-methods, partners, rules", "PASS"),
            ("Inventory", "dashboard, units, categories, medicines, items, suppliers, locations, batches, purchase-orders, purchase-invoices, purchase-returns, supplier-dues, transfers, adjustments", "PASS render; several data APIs fail"),
            ("Laboratory", "/admin/laboratory/labTests, /admin/laboratory/packages", "PASS"),
            ("Monitoring", "/admin/monitoring/reports, /admin/monitoring/audit", "PASS"),
            ("Administration", "/admin/roles, /admin/settings/organization, notifications, system, security", "PASS render"),
        ],
        [1.25, 4.25, 1.0],
    )

    doc.add_heading("CRUD and State-Change Matrix", level=1)
    add_table(doc, ["Section", "Operation Tested", "Status", "Evidence / Notes"], status_rows(), [1.55, 1.75, 0.75, 2.45])

    doc.add_heading("Confirmed Blockers", level=1)
    add_table(
        doc,
        ["Priority", "Area", "Observed Failure", "Recommended Fix"],
        [
            ("P1", "Patient creation", "POST /api/admin/patients returns 500 because AdminPatientsController.Create starts a transaction. MongoDB EF Core provider does not support transactions.", "Remove transaction usage for Mongo provider or replace with provider-supported consistency pattern."),
            ("P1", "Billing list pages", "Mongo cannot translate helper method projections such as MapChargeDefinition and MapPaymentMethod inside IQueryable Select.", "Materialize with ToListAsync first, then map in memory; or project directly to DTO fields without helper calls inside the query."),
            ("P1", "Inventory read/write return paths", "Mongo cannot translate several Include/OrderBy query shapes in EfInventoryAdminService for medicines/items/locations/purchase invoices/returns/transfers.", "Replace Include-heavy EF queries with explicit materialization and manual joins/lookups compatible with Mongo EF."),
            ("P2", "Settings API contract", "Routes for notifications/system/security render, but standalone GET /api/settings/notifications, /system, /security return 404. The app appears to load these from /api/settings/control-center.", "Either add GET endpoints matching PUT endpoints or document and enforce control-center as the read source."),
            ("P2", "Delete semantics", "Only hospital services expose hard DELETE. Most modules use isActive/status/lifecycle closure.", "Keep this intentionally soft-delete model, but label UI actions as deactivate/archive where applicable."),
        ],
        [0.6, 1.3, 2.4, 2.2],
    )

    doc.add_heading("Testing Notes", level=1)
    add_bullets(doc, [
        "Backend was restarted because the existing AngularApp4 process caused auth requests to hang. After restart, login succeeded directly and through Angular.",
        "CRUD test stamp: 9086381. Test-created active resources were deactivated where the API supports status changes. Hospital service test data was hard-deleted.",
        "Because patient creation is blocked by a real backend exception, dependent create flows for admissions and invoices were not forced against existing patient records.",
        "Because inventory location/medicine/item creation hit provider translation failures, dependent purchase order, transfer, and adjustment create flows were blocked.",
    ])

    doc.add_heading("Next Fix Order", level=1)
    add_bullets(doc, [
        "Fix Mongo provider incompatibilities in patient creation first; it unlocks patients, admissions, and billing invoice CRUD.",
        "Fix DTO mapping/projection patterns in billing and inventory list endpoints; these currently break rendered pages that depend on those API responses.",
        "Add or document settings read endpoints so each settings route has a clear read/update contract.",
        "Re-run the same matrix after fixes and update the pass/fail counts.",
    ])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
